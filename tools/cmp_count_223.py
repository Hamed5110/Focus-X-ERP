import sys
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

CUBE = r"C:\Users\Hamed Ali Khan\Downloads\Trade_Receivables_180_60810_Project_Tracking_III_Report_Atlas60810_ATLAS_ALUMINUM_W_L_L_.xlsx"
SQLCSV = r"C:\Users\Hamed Ali Khan\Focus-X-ERP\tools\acct_level_out.csv"

wb = load_workbook(CUBE, data_only=True)
ws = wb.active
headers = [str(ws.cell(6, c).value or "").replace("\xa0", " ").strip() for c in range(1, ws.max_column + 1)]
idx = {h: i for i, h in enumerate(headers, 1)}
name_col = 1
if "Particulars" in idx:
    name_col = idx["Particulars"]
cube_names = []
for r in range(7, ws.max_row + 1):
    nm = str(ws.cell(r, name_col).value or "").strip()
    if not nm or "grand" in nm.lower() or nm.lower() == "total":
        continue
    cube_names.append(nm)
wb.close()

import csv
sql_rows = []
with open(SQLCSV, encoding="utf-8-sig") as f:
    for row in csv.DictReader(f):
        sql_rows.append(row)

sql3 = [r for r in sql_rows if r["ReportStatus"] == "3"]
sql_names = [r["Name"].strip() for r in sql3]

print("cube rows:", len(cube_names), "distinct:", len(set(cube_names)))
print("sql status3 rows:", len(sql3), "distinct names:", len(set(sql_names)))

cs, ss = set(cube_names), set(sql_names)
print("\nSQL-only (in SQL, not in cube):")
for n in sorted(ss - cs):
    r = [x for x in sql3 if x["Name"].strip() == n][0]
    print(f"  {n[:60]:62} C={r['Total Contract Amount']:>12} A={r['Adv. Rct Amount']:>12}")
print("\nCube-only (in cube, not in SQL):")
for n in sorted(cs - ss):
    print(" ", n[:70])
