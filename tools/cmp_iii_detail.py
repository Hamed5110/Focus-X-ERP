# -*- coding: utf-8 -*-
import sys
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

CUBE = r"C:\Users\Hamed Ali Khan\Downloads\Trade_Receivables_180_931647_Project_Tracking_III_Report_Atlas931647_ATLAS_ALUMINUM_W_L_L_.xlsx"
SQLX = r"C:\Users\Hamed Ali Khan\Downloads\III_testing480231_ATLAS_ALUMINUM_W_L_L_.xlsx"


def num(v):
    if v is None or v is False:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "")
    if s in ("", "-", "None", "False"):
        return 0.0
    try:
        return float(s)
    except ValueError:
        return 0.0


def load(path):
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    headers = [
        str(ws.cell(6, c).value or "").replace("\xa0", " ").strip()
        for c in range(1, ws.max_column + 1)
    ]
    idx = {h: i for i, h in enumerate(headers, 1)}
    data = {}
    tot = None
    for r in range(7, ws.max_row + 1):
        name = str(ws.cell(r, 1).value or "").strip()
        code = str(ws.cell(r, idx["Code"]).value or "").strip() if "Code" in idx else ""
        rec = {
            "name": name,
            "code": code,
            "c": num(ws.cell(r, idx["Total Contract Amount"]).value),
            "a": num(ws.cell(r, idx["Adv. Rct Amount"]).value),
        }
        if "grand" in name.lower() or name.lower() == "total":
            tot = rec
            continue
        if not code:
            continue
        data[code] = rec
    wb.close()
    return data, tot


cube, ct = load(CUBE)
sqlx, st = load(SQLX)
common = set(cube) & set(sqlx)

print("=== ALL Contract |d|>0.01 ===")
nc = 0
for k in sorted(common):
    d = sqlx[k]["c"] - cube[k]["c"]
    if abs(d) > 0.01:
        nc += 1
        print(f"{k:10} cube={cube[k]['c']:.2f} sql={sqlx[k]['c']:.2f} d={d:.2f}")
print("count", nc)

print("\n=== ALL Adv |d|>0.01 ===")
na = 0
for k in sorted(common):
    d = sqlx[k]["a"] - cube[k]["a"]
    if abs(d) > 0.01:
        na += 1
        print(f"{k:10} cube={cube[k]['a']:.2f} sql={sqlx[k]['a']:.2f} d={d:.2f}")
print("count", na)

print("\n=== After ROUND(sql,0) vs cube ===")
mc = ma = 0
for k in sorted(common):
    if abs(round(sqlx[k]["c"]) - cube[k]["c"]) > 0.01:
        mc += 1
        print(f"C {k} cube={cube[k]['c']} r={round(sqlx[k]['c'])} raw={sqlx[k]['c']}")
    if abs(round(sqlx[k]["a"]) - cube[k]["a"]) > 0.01:
        ma += 1
        print(f"A {k} cube={cube[k]['a']} r={round(sqlx[k]['a'])} raw={sqlx[k]['a']}")
print("fail contract", mc, "adv", ma)

print("\n=== Totals ===")
print("cube rows", sum(v["c"] for v in cube.values()), sum(v["a"] for v in cube.values()))
print("sql  rows", sum(v["c"] for v in sqlx.values()), sum(v["a"] for v in sqlx.values()))
print("cube footer", ct)
print("sql  footer", st)
print(
    "FOOTER TRAP: cube grand Contract",
    ct["c"] if ct else None,
    "vs row-sum",
    sum(v["c"] for v in cube.values()),
    "delta",
    (ct["c"] - sum(v["c"] for v in cube.values())) if ct else None,
)
