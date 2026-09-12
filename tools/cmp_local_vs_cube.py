import csv
import sys
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

LOCALCSV = r"C:\Users\Hamed Ali Khan\Focus-X-ERP\tools\diag_status3.csv"
CUBE = r"C:\Users\Hamed Ali Khan\Downloads\Trade_Receivables_180_821618_Project_Tracking_III_Report_Atlas821618_ATLAS_ALUMINUM_W_L_L_.xlsx"


def num(v):
    if v is None or v is False:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "")
    if s in ("", "-", "None", "False"):
        return 0.0
    try:
        return float(s)
    except ValueError:
        return 0.0


local = {}
with open(LOCALCSV, newline="", encoding="utf-8-sig") as f:
    for row in csv.DictReader(f):
        name = row["Report Status"].strip()
        local[name] = {"mid": row["No. of Accounts"].strip(),
                       "c": num(row["Total Contract Amount"]),
                       "a": num(row["Adv. Rct Amount"])}

wb = load_workbook(CUBE, data_only=True)
ws = wb.active
hr = None
for r in range(1, 12):
    vals = [str(ws.cell(r, c).value or "").replace("\xa0", " ").strip() for c in range(1, ws.max_column + 1)]
    if "Total Contract Amount" in vals:
        hr = r
        headers = vals
        break
idx = {h: i for i, h in enumerate(headers, 1)}
cube = {}
for r in range(hr + 1, ws.max_row + 1):
    name = str(ws.cell(r, idx["Particulars"]).value or "").strip()
    if not name or "grand" in name.lower() or name.lower() == "total":
        continue
    cube[name] = {"code": str(ws.cell(r, idx["Code"]).value or "").strip(),
                  "c": num(ws.cell(r, idx["Total Contract Amount"]).value),
                  "a": num(ws.cell(r, idx["Adv. Rct Amount"]).value),
                  "so_date": str(ws.cell(r, idx["Sales Order Date"]).value or "").strip(),
                  "site": str(ws.cell(r, idx["Site Status"]).value or "").strip(),
                  "pipe": str(ws.cell(r, idx["Pipeline"]).value or "").strip()}
wb.close()

print("Local query rows:", len(local), " Cube rows:", len(cube))
ls, cs = set(local), set(cube)

print("\n==== In LOCAL query but NOT in cube (cube excludes) ====")
for n in sorted(ls - cs):
    v = local[n]
    print(f"  id={v['mid']:<7} {n[:55]:57} C={v['c']:>10.2f} A={v['a']:>10.2f}")

print("\n==== In CUBE but NOT in local query (query misses) ====")
for n in sorted(cs - ls):
    v = cube[n]
    print(f"  code={v['code']:<10} {n[:45]:47} C={v['c']:>10.2f} A={v['a']:>10.2f} SODate={v['so_date']:12} Site={v['site']:10} Pipe={v['pipe']}")
