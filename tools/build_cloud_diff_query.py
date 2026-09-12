import sys, json
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

DL = r"C:\Users\Hamed Ali Khan\Downloads"
f3 = DL + r"\Trade_Receivables_180_145256_Project_Tracking_III_Report_Atlas145256_ATLAS_ALUMINUM_W_L_L_.xlsx"

def num(v):
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace(",", "").replace("\xa0", "").strip()
    try:
        return float(s)
    except ValueError:
        return 0.0

wb = load_workbook(f3, data_only=True)
ws = wb[wb.sheetnames[0]]
rows = [[("" if v is None else v) for v in r] for r in ws.iter_rows(values_only=True)]
wb.close()

hi = next(i for i, r in enumerate(rows) if str(r[0]).strip() == "Particulars")
hdr = [str(x).strip() for x in rows[hi]]
ci_code = hdr.index("Code")
ci_c = hdr.index("Total Contract Amount")
ci_a = hdr.index("Adv. Rct Amount")

vals = []
for r in rows[hi + 1:]:
    name = str(r[0]).strip()
    if name in ("", "Grand Total"):
        continue
    code = str(r[ci_code]).strip().replace("'", "''")
    vals.append((code, num(r[ci_c]), num(r[ci_a])))

print("cube III rows:", len(vals))
values_sql = ",\n".join(f"('{c}', {cc}, {aa})" for c, cc, aa in vals)

# detail inner query = the summary's inner subquery x, plus Code exposed
inner = open(r"C:\Users\Hamed Ali Khan\Focus-X-ERP\reports\sql\III Summary Report.sql", encoding="utf-8").read()
start = inner.index("    SELECT\n        acc.ReportStatus,")
end = inner.index(") x\nGROUP BY x.ReportStatus")
detail = inner[start:end]
# expose the account code: acc subquery has iMasterId; add sCode via scalar lookup
detail = detail.replace(
    "acc.sName AS Name,",
    "acc.sName AS Name,\n        acc.iMasterId,\n        (SELECT sCode FROM dbo.mCore_Account mc WHERE mc.iMasterId = acc.iMasterId) AS Code,",
    1,
)

query = f"""SELECT
    ISNULL(s.Code, c.Code) AS Code,
    s.Name AS SqlName,
    CAST(ISNULL(s.[Total Contract Amount], 0) AS decimal(18,0)) AS SqlContract,
    CAST(c.CubeContract AS decimal(18,0)) AS CubeContract,
    CAST(ISNULL(s.[Adv. Rct Amount], 0) AS decimal(18,0)) AS SqlAdv,
    CAST(c.CubeAdv AS decimal(18,0)) AS CubeAdv,
    CAST(ISNULL(s.[Total Contract Amount], 0) - c.CubeContract AS decimal(18,0)) AS DiffContract,
    CAST(ISNULL(s.[Adv. Rct Amount], 0) - c.CubeAdv AS decimal(18,0)) AS DiffAdv,
    CASE WHEN s.Code IS NULL THEN 'CUBE ONLY'
         WHEN c.Code IS NULL THEN 'SQL ONLY'
         ELSE 'AMOUNT DIFF' END AS Issue
FROM (
{detail}
) s
FULL OUTER JOIN (
    VALUES
{values_sql}
) AS c(Code, CubeContract, CubeAdv) ON c.Code = s.Code AND s.ReportStatus = 3
WHERE s.ReportStatus = 3 OR s.Code IS NULL
  AND 1 = 1
ORDER BY Issue, Code
"""

# hmm: FULL JOIN + WHERE s.ReportStatus=3 would kill CUBE ONLY rows; fix predicate
query = query.replace(
    "WHERE s.ReportStatus = 3 OR s.Code IS NULL\n  AND 1 = 1",
    "WHERE (s.ReportStatus = 3 OR s.Code IS NULL)\n  AND (s.Code IS NULL OR c.Code IS NULL\n       OR ISNULL(s.[Total Contract Amount],0) <> c.CubeContract\n       OR ISNULL(s.[Adv. Rct Amount],0) <> c.CubeAdv)",
)

with open(r"C:\Users\Hamed Ali Khan\Focus-X-ERP\tools\cloud_diff_query.sql", "w", encoding="utf-8") as f:
    f.write(query)
print("query chars:", len(query))
print("saved to tools/cloud_diff_query.sql")
