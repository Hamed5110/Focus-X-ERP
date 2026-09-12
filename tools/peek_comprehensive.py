import sys
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

base = r"C:\Users\Hamed Ali Khan\Downloads"
files = [
    base + r"\Trade_Receivables_180_643162_Project_Tracking_I_Report_Atlas643162_ATLAS_ALUMINUM_W_L_L_.xlsx",
    base + r"\Trade_Receivables_180_264376_Project_Tracking_II_Report_Atlas264376_ATLAS_ALUMINUM_W_L_L_.xlsx",
    base + r"\Trade_Receivables_180_145256_Project_Tracking_III_Report_Atlas145256_ATLAS_ALUMINUM_W_L_L_.xlsx",
    base + r"\Comprehensive_Project_Tracking_Report779253_ATLAS_ALUMINUM_W_L_L_.xlsx",
]

for f in files:
    print("=" * 100)
    print(f.split("\\")[-1])
    wb = load_workbook(f, data_only=True)
    print("sheets:", wb.sheetnames)
    ws = wb.active
    print("dims:", ws.max_row, "x", ws.max_column)
    for r in range(1, min(8, ws.max_row + 1)):
        vals = [str(ws.cell(r, c).value or "").replace("\xa0", " ").strip()[:26] for c in range(1, min(ws.max_column + 1, 18))]
        print(" ", r, vals)
    wb.close()
