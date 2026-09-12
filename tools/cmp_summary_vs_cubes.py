import sys
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

base = r"C:\Users\Hamed Ali Khan\Downloads"
cube_files = {
    "1 Pending - I": base + r"\Trade_Receivables_180_643162_Project_Tracking_I_Report_Atlas643162_ATLAS_ALUMINUM_W_L_L_.xlsx",
    "2 In Progress - II": base + r"\Trade_Receivables_180_264376_Project_Tracking_II_Report_Atlas264376_ATLAS_ALUMINUM_W_L_L_.xlsx",
    "3 Partial Consumed - III": base + r"\Trade_Receivables_180_145256_Project_Tracking_III_Report_Atlas145256_ATLAS_ALUMINUM_W_L_L_.xlsx",
}
comp = base + r"\Comprehensive_Project_Tracking_Report779253_ATLAS_ALUMINUM_W_L_L_.xlsx"


def load_cube(path):
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    hr = None
    headers = None
    for r in range(1, 12):
        vals = [str(ws.cell(r, c).value or "").replace("\xa0", " ").strip() for c in range(1, ws.max_column + 1)]
        if "Total Contract Amount" in vals:
            hr = r
            headers = vals
            break
    idx = {h: i for i, h in enumerate(headers, 1) if h}
    rows = {}
    for r in range(hr + 1, ws.max_row + 1):
        n = str(ws.cell(r, idx["Particulars"]).value or "").replace("\xa0", " ").strip()
        if n and "grand" not in n.lower() and n.lower() != "total":
            rows[n] = {
                "C": float(ws.cell(r, idx["Total Contract Amount"]).value or 0),
                "A": float(ws.cell(r, idx["Adv. Rct Amount"]).value or 0),
                "P": float(ws.cell(r, idx["Plan Value"]).value or 0) if "Plan Value" in idx else 0.0,
            }
    wb.close()
    return rows


print("=== Comprehensive Project Tracking Report (full) ===")
wb = load_workbook(comp, data_only=True)
ws = wb.active
comp_rows = {}
for r in range(7, ws.max_row + 1):
    vals = [ws.cell(r, c).value for c in range(1, 7)]
    if vals[0]:
        print(" ", vals)
        comp_rows[str(vals[0]).strip()] = vals[1:]
wb.close()

print()
print("=== Cube I/II/III totals vs Comprehensive ===")
print(f"{'Status':28} {'CubeC':>12} {'CompC':>12} {'dC':>9} {'CubeA':>12} {'CompA':>12} {'dA':>9} {'CubeN':>6} {'CompN':>6}")
for status, f in cube_files.items():
    rows = load_cube(f)
    tc = sum(r["C"] for r in rows.values())
    ta = sum(r["A"] for r in rows.values())
    n = len(rows)
    cr = comp_rows.get(status)
    if cr:
        cc, ca, cb, cp, cn = [float(x or 0) for x in cr]
        print(f"{status:28} {tc:>12,.0f} {cc:>12,.0f} {tc-cc:>9,.0f} {ta:>12,.0f} {ca:>12,.0f} {ta-ca:>9,.0f} {n:>6} {cn:>6,.0f}")
    else:
        print(f"{status:28} {tc:>12,.0f} {'-':>12} {'':>9} {ta:>12,.0f} {'-':>12} {'':>9} {n:>6}")
