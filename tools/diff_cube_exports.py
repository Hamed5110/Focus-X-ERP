import sys
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

A = r"C:\Users\Hamed Ali Khan\Downloads\Trade_Receivables_180_60810_Project_Tracking_III_Report_Atlas60810_ATLAS_ALUMINUM_W_L_L_.xlsx"
B = r"C:\Users\Hamed Ali Khan\Downloads\Trade_Receivables_180_821618_Project_Tracking_III_Report_Atlas821618_ATLAS_ALUMINUM_W_L_L_.xlsx"


def num(v):
    if v is None or v is False:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "")
    try:
        return float(s)
    except ValueError:
        return 0.0


def load(path):
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    hr = None
    for r in range(1, 12):
        vals = [str(ws.cell(r, c).value or "").replace("\xa0", " ").strip() for c in range(1, ws.max_column + 1)]
        if "Total Contract Amount" in vals:
            hr = r
            headers = vals
            break
    idx = {h: i for i, h in enumerate(headers, 1)}
    rows = {}
    for r in range(hr + 1, ws.max_row + 1):
        n = str(ws.cell(r, idx["Particulars"]).value or "").strip()
        if n and "grand" not in n.lower() and n.lower() != "total":
            rows[n] = {h: ws.cell(r, i).value for h, i in idx.items()}
    wb.close()
    return rows


ra, rb = load(A), load(B)
print(f"60810: {len(ra)} rows | 821618: {len(rb)} rows")

only_a = set(ra) - set(rb)
only_b = set(rb) - set(ra)
print("\nOnly in 60810:", sorted(only_a))
print("\nOnly in 821618:", sorted(only_b))

print("\nValue differences on common names:")
diffs = 0
for n in sorted(set(ra) & set(rb)):
    rowdiffs = []
    for h in ra[n]:
        va, vb = ra[n][h], rb[n][h]
        if isinstance(ra[n][h], (int, float)) or isinstance(vb, (int, float)):
            if abs(num(va) - num(vb)) > 0.005:
                rowdiffs.append(f"{h}: {va} -> {vb}")
        elif str(va or "").strip() != str(vb or "").strip():
            rowdiffs.append(f"{h}: '{va}' -> '{vb}'")
    if rowdiffs:
        diffs += 1
        print(f"  {n}: " + " | ".join(rowdiffs))
print(f"\nTotal rows with any difference: {diffs}")
