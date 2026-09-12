import sys, glob
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

base = r"C:\Users\Hamed Ali Khan\Downloads"
f1 = glob.glob(base + r"\*Project_Tracking_I_Report_Atlas*.xlsx")[0]
f2 = glob.glob(base + r"\*Project_Tracking_II_Report_Atlas*.xlsx")[0]
f3 = glob.glob(base + r"\*Project_Tracking_III_Report_Atlas821618*.xlsx")[0]


def load(path):
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    hr = None
    headers = None
    for r in range(1, 12):
        vals = [str(ws.cell(r, c).value or "").replace("\xa0", " ").strip() for c in range(1, ws.max_column + 1)]
        if "Total Contract Amount" in vals:
            hr = r
            headers = vals
            break
    if hr is None:
        wb.close()
        return {}
    idx = {h: i for i, h in enumerate(headers, 1)}
    rows = {}
    for r in range(hr + 1, ws.max_row + 1):
        n = str(ws.cell(r, idx["Particulars"]).value or "").strip()
        if n and "grand" not in n.lower() and n.lower() != "total":
            rows[n] = True
    wb.close()
    return rows


n1 = load(f1)
n2 = load(f2)
n3 = load(f3)

print("Cube I   rows:", len(n1))
print("Cube II  rows:", len(n2))
print("Cube III rows:", len(n3))
print()
print("Names in BOTH I and III:", sorted(set(n1) & set(n3)))
print("Names in BOTH II and III:", sorted(set(n2) & set(n3)))
print("Names in BOTH I and II:", sorted(set(n1) & set(n2)))
