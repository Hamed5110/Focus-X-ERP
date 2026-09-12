import sys
from openpyxl import load_workbook
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
path = r"C:\Users\Hamed Ali Khan\Downloads\Monthly_Commission_Atlas901133_ATLAS_ALUMINUM_W_L_L_.xlsx"
wb = load_workbook(path, data_only=True)
ws = wb[wb.sheetnames[0]]
print("sheets", wb.sheetnames, "rows", ws.max_row, "cols", ws.max_column)
for i, r in enumerate(ws.iter_rows(values_only=True), 1):
    if i > 12:
        break
    cells = [("" if v is None else str(v).replace("\n"," "))[:55] for v in r]
    while cells and cells[-1] == "":
        cells.pop()
    print(f"{i:03d}| " + " || ".join(cells))
wb.close()
