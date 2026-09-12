import sys
from collections import Counter
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

SQLX = r"C:\Users\Hamed Ali Khan\Downloads\iii_Project_Summ_SQL_report531004_ATLAS_ALUMINUM_W_L_L_.xlsx"
CUBE = r"C:\Users\Hamed Ali Khan\Downloads\Trade_Receivables_180_821618_Project_Tracking_III_Report_Atlas821618_ATLAS_ALUMINUM_W_L_L_.xlsx"


def num(v):
    if v is None or v is False:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "")
    if s in ("", "-", "None", "False"):
        return 0.0
    try:
        return float(s)
    except ValueError:
        return 0.0


def find_header_row(ws, must_have):
    for r in range(1, 12):
        vals = [str(ws.cell(r, c).value or "").replace("\xa0", " ").strip() for c in range(1, ws.max_column + 1)]
        if all(any(mh == v for v in vals) for mh in must_have):
            return r, vals
    return None, None


wb = load_workbook(SQLX, data_only=True)
ws = wb.active
hr, headers = find_header_row(ws, ["Report Status", "No. of Accounts"])
idx = {h: i for i, h in enumerate(headers, 1)}
sql_list = []
for r in range(hr + 1, ws.max_row + 1):
    name = str(ws.cell(r, idx["Report Status"]).value or "").strip()
    if not name or "grand" in name.lower() or name.lower() == "total":
        continue
    sql_list.append({
        "name": name,
        "mid": int(num(ws.cell(r, idx["No. of Accounts"]).value)),
        "c": num(ws.cell(r, idx["Total Contract Amount"]).value),
        "a": num(ws.cell(r, idx["Adv. Rct Amount"]).value),
        "p": num(ws.cell(r, idx["Plan Value"]).value),
    })
wb.close()

wb = load_workbook(CUBE, data_only=True)
ws = wb.active
hr, headers = find_header_row(ws, ["Total Contract Amount"])
idx = {h: i for i, h in enumerate(headers, 1)}
cube_list = []
for r in range(hr + 1, ws.max_row + 1):
    name = str(ws.cell(r, idx["Particulars"]).value or "").strip()
    if not name or "grand" in name.lower() or name.lower() == "total":
        continue
    cube_list.append({
        "name": name,
        "code": str(ws.cell(r, idx["Code"]).value or "").strip(),
        "c": num(ws.cell(r, idx["Total Contract Amount"]).value),
        "a": num(ws.cell(r, idx["Adv. Rct Amount"]).value),
        "p": num(ws.cell(r, idx["Plan Value"]).value),
    })
wb.close()

print("SQL rows:", len(sql_list), " CUBE rows:", len(cube_list))

dup_sql = [n for n, k in Counter(x["name"] for x in sql_list).items() if k > 1]
dup_cube = [n for n, k in Counter(x["name"] for x in cube_list).items() if k > 1]
print("\nDuplicate names in SQL export:", dup_sql)
print("Duplicate names in cube export:", dup_cube)

cube_by_name = {x["name"]: x for x in cube_list}
print("\n==== Amount mismatches on common names (|d|>0.5) ====")
for s in sql_list:
    c = cube_by_name.get(s["name"])
    if not c:
        continue
    dc, da = s["c"] - c["c"], s["a"] - c["a"]
    if abs(dc) > 0.5 or abs(da) > 0.5:
        print(f"  id={s['mid']:<7} {s['name'][:45]:47} SQL C={s['c']:>10.2f} A={s['a']:>10.2f} | CUBE C={c['c']:>10.2f} A={c['a']:>10.2f} (code {c['code']})")

print("\n==== Fuzzy look for the 2 extras in cube ====")
for probe in ("Ahmed Abdulla", "Hussain Ali", "Jaffar", "Hussain"):
    hits = [x for x in cube_list if probe.lower() in x["name"].lower()]
    for h in hits:
        print(f"  probe={probe:16} cube: {h['name'][:55]:57} code={h['code']:>10} C={h['c']:>10.2f} A={h['a']:>10.2f}")

print("\nSQL totals: C={:.2f} A={:.2f} P={:.2f}".format(sum(x["c"] for x in sql_list), sum(x["a"] for x in sql_list), sum(x["p"] for x in sql_list)))
print("CUBE totals: C={:.2f} A={:.2f} P={:.2f}".format(sum(x["c"] for x in cube_list), sum(x["a"] for x in cube_list), sum(x["p"] for x in cube_list)))
