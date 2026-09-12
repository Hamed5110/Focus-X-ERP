import openpyxl, glob

base = r"C:\Users\Hamed Ali Khan\Downloads"
f3 = glob.glob(base + r"\*Project_Tracking_III_Report_Atlas821618*.xlsx")[0]
print("file:", f3)
wb = openpyxl.load_workbook(f3, read_only=True, data_only=True)
print("sheets:", wb.sheetnames)
ws = wb.active
for i, r in enumerate(ws.iter_rows(values_only=True)):
    print(i, [str(c)[:28] if c is not None else None for c in r][:12])
    if i >= 8:
        break
