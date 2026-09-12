import sys
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

FILES = {
    "III (fresh 821618)": r"C:\Users\Hamed Ali Khan\Downloads\Trade_Receivables_180_821618_Project_Tracking_III_Report_Atlas821618_ATLAS_ALUMINUM_W_L_L_.xlsx",
    "III (older 60810)": r"C:\Users\Hamed Ali Khan\Downloads\Trade_Receivables_180_60810_Project_Tracking_III_Report_Atlas60810_ATLAS_ALUMINUM_W_L_L_.xlsx",
    "II (314796)": r"C:\Users\Hamed Ali Khan\Downloads\Trade_Receivables_180_314796_Project_Tracking_II_Report_Atlas314796_ATLAS_ALUMINUM_W_L_L_.xlsx",
    "I (362696)": r"C:\Users\Hamed Ali Khan\Downloads\Trade_Receivables_180_362696_Project_Tracking_I_Report_Atlas362696_ATLAS_ALUMINUM_W_L_L_.xlsx",
}

PROBES = ["Dr. Muneer Mahdi", "Mr. Ahmed Abdulla Jaffar", "Mr. Hussain Ali",
          "Bayan Naser Salman", "Mr. Jaafar Jasim Ebrahim Hasan Ali", "Mr. Mohammed Jawad",
          "Mrs. Zahra Al Shaikh", "Zuhair Abbas Ahmed"]

for label, path in FILES.items():
    try:
        wb = load_workbook(path, data_only=True)
    except Exception as e:
        print(f"{label}: cannot open ({e})")
        continue
    ws = wb.active
    hr = None
    for r in range(1, 12):
        vals = [str(ws.cell(r, c).value or "").replace("\xa0", " ").strip() for c in range(1, ws.max_column + 1)]
        if "Total Contract Amount" in vals:
            hr = r
            headers = vals
            break
    if hr is None:
        print(f"{label}: no header found")
        wb.close()
        continue
    idx = {h: i for i, h in enumerate(headers, 1)}
    ncol = idx.get("Particulars", 1)
    names = set()
    for r in range(hr + 1, ws.max_row + 1):
        n = str(ws.cell(r, ncol).value or "").strip()
        if n and "grand" not in n.lower() and n.lower() != "total":
            names.add(n)
    wb.close()
    print(f"\n== {label}: {len(names)} rows ==")
    for p in PROBES:
        print(f"  {'IN ' if p in names else 'OUT'}  {p}")
