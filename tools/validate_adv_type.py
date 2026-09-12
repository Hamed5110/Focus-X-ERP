# -*- coding: utf-8 -*-
import csv
import sys
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8", errors="replace")


def num(v):
    if v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", ""))
    except Exception:
        return 0.0


wb = load_workbook(
    r"C:\Users\Hamed Ali Khan\Downloads\Trade_Receivables_180_931647_Project_Tracking_III_Report_Atlas931647_ATLAS_ALUMINUM_W_L_L_.xlsx",
    data_only=True,
)
ws = wb.active
headers = [
    str(ws.cell(6, c).value or "").replace("\xa0", " ").strip()
    for c in range(1, ws.max_column + 1)
]
idx = {h: i for i, h in enumerate(headers, 1)}
cube = {}
for r in range(7, ws.max_row + 1):
    name = str(ws.cell(r, 1).value or "").strip()
    if "grand" in name.lower():
        continue
    code = str(ws.cell(r, idx["Code"]).value or "").strip()
    if not code:
        continue
    cube[code] = {
        "c": num(ws.cell(r, idx["Total Contract Amount"]).value),
        "a": num(ws.cell(r, idx["Adv. Rct Amount"]).value),
    }
wb.close()

sql = {}
with open(
    r"C:\Users\Hamed Ali Khan\Focus-X-ERP\tools\sql_adv_type_round.csv",
    encoding="utf-8-sig",
) as f:
    for row in csv.DictReader(f):
        sql[row["sCode"]] = {k: num(row[k]) for k in ("AdvRaw", "AdvRoundT", "AdvRoundType")}

# also load contract rounds
with open(
    r"C:\Users\Hamed Ali Khan\Focus-X-ERP\tools\sql_round_name.csv",
    encoding="utf-8-sig",
) as f:
    for row in csv.DictReader(f):
        sql.setdefault(row["sCode"], {})
        sql[row["sCode"]]["ContractRaw"] = num(row["ContractRaw"])
        sql[row["sCode"]]["ContractRoundT"] = num(row["ContractRoundT"])

common = set(cube) & set(sql)
for label in ("AdvRoundT", "AdvRoundType"):
    fails = [
        (c, cube[c]["a"], sql[c]["AdvRaw"], sql[c][label])
        for c in sorted(common)
        if abs(sql[c][label] - cube[c]["a"]) > 0.01
    ]
    print(label, "fails", len(fails), "sum", sum(sql[c][label] for c in common))
    for x in fails[:30]:
        print(x)

cfails = [
    (c, cube[c]["c"], sql[c]["ContractRaw"], sql[c]["ContractRoundT"])
    for c in sorted(common)
    if abs(sql[c]["ContractRoundT"] - cube[c]["c"]) > 0.01
]
print("\nContractRoundT fails", len(cfails), "sum", sum(sql[c]["ContractRoundT"] for c in common))
print("cube c/a", sum(v["c"] for v in cube.values()), sum(v["a"] for v in cube.values()))
