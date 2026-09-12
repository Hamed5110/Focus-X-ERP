import sys
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

CUBE = r"C:\Users\Hamed Ali Khan\Downloads\Trade_Receivables_180_931647_Project_Tracking_III_Report_Atlas931647_ATLAS_ALUMINUM_W_L_L_.xlsx"
SQLX = r"C:\Users\Hamed Ali Khan\Downloads\III_testing480231_ATLAS_ALUMINUM_W_L_L_.xlsx"


def num(v):
    if v is None or v is False:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "")
    if s in ("", "-", "None", "False"):
        return 0.0
    try:
        return float(s)
    except ValueError:
        return 0.0


def load(path):
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    headers = [str(ws.cell(6, c).value or "").replace("\xa0", " ").strip() for c in range(1, ws.max_column + 1)]
    print("\nFILE", path.split("\\")[-1])
    print("headers", headers)
    idx = {h: i for i, h in enumerate(headers, 1)}
    pn = "Production Note Amount" if "Production Note Amount" in idx else "Production Note Amount "
    data = {}
    tot = None
    for r in range(7, ws.max_row + 1):
        name = str(ws.cell(r, 1).value or "").strip()
        code = str(ws.cell(r, idx["Code"]).value or "").strip() if "Code" in idx else ""
        rec = {
            "name": name,
            "code": code,
            "c": num(ws.cell(r, idx["Total Contract Amount"]).value),
            "a": num(ws.cell(r, idx["Adv. Rct Amount"]).value),
            "b": num(ws.cell(r, idx["Balance Amount"]).value),
            "s": num(ws.cell(r, idx["Sales Job Order"]).value),
            "p": num(ws.cell(r, idx[pn]).value) if pn in idx else 0.0,
        }
        if "grand" in name.lower() or name.lower() == "total":
            tot = rec
            continue
        if not code:
            continue
        data[code] = rec
    wb.close()
    sm = {k: round(sum(v[k] for v in data.values()), 2) for k in "cabsp"}
    print("n", len(data), "grand", tot)
    print("sum rows", sm)
    return data, tot


cube, ct = load(CUBE)
sqlx, st = load(SQLX)
cc, sc = set(cube), set(sqlx)
print("\nonly cube", sorted(cc - sc))
print("only sql", sorted(sc - cc))

print("\n==== CONTRACT/ADV material |d|>0.51 ====")
rows = []
for code in sorted(cc & sc):
    c, s = cube[code], sqlx[code]
    dc, da = s["c"] - c["c"], s["a"] - c["a"]
    if abs(dc) > 0.51 or abs(da) > 0.51:
        rows.append((abs(dc) + abs(da), code, c, s, dc, da))
print("n material", len(rows))
rows.sort(reverse=True)
for mag, code, c, s, dc, da in rows[:40]:
    print(f"\n{code} {c['name'][:45]}")
    if abs(dc) > 0.51:
        print(f"  CONTRACT cube={c['c']:.2f} sql={s['c']:.2f} d={dc:.2f}")
    if abs(da) > 0.51:
        print(f"  ADV      cube={c['a']:.2f} sql={s['a']:.2f} d={da:.2f}")

print("\nCUBE sum c/a", round(sum(v["c"] for v in cube.values()), 2), round(sum(v["a"] for v in cube.values()), 2))
print("SQL  sum c/a", round(sum(v["c"] for v in sqlx.values()), 2), round(sum(v["a"] for v in sqlx.values()), 2))
print("common cube c/a", round(sum(cube[k]["c"] for k in cc & sc), 2), round(sum(cube[k]["a"] for k in cc & sc), 2))
print("common sql  c/a", round(sum(sqlx[k]["c"] for k in cc & sc), 2), round(sum(sqlx[k]["a"] for k in cc & sc), 2))
