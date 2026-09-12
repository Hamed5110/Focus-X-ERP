import sys
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

path = r"C:\Users\Hamed Ali Khan\Downloads\Monthly_Commission_Atlas363679_ATLAS_ALUMINUM_W_L_L_.xlsx"
wb = load_workbook(path, data_only=True)
print("sheets:", wb.sheetnames)
for sn in wb.sheetnames:
    ws = wb[sn]
    print(f"\n--- {sn} rows={ws.max_row} cols={ws.max_column} ---")
    for i, r in enumerate(ws.iter_rows(values_only=True), 1):
        if i > 80:
            print("...truncated...")
            break
        if any(v is not None and str(v).strip() != "" for v in r):
            cells = [("" if v is None else str(v).replace("\n", " | "))[:60] for v in r]
            while cells and cells[-1] == "":
                cells.pop()
            print(f"{i:03d}| " + " || ".join(cells))
wb.close()
