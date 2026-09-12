import sys
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

for label, path in [
    ("FRESH 821618", r"C:\Users\Hamed Ali Khan\Downloads\Trade_Receivables_180_821618_Project_Tracking_III_Report_Atlas821618_ATLAS_ALUMINUM_W_L_L_.xlsx"),
    ("SQL 531004", r"C:\Users\Hamed Ali Khan\Downloads\iii_Project_Summ_SQL_report531004_ATLAS_ALUMINUM_W_L_L_.xlsx"),
]:
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    print(f"\n===== {label}: first 6 rows x first 12 cols =====")
    for r in range(1, 7):
        vals = []
        for c in range(1, min(ws.max_column, 12) + 1):
            v = ws.cell(r, c).value
            if v is not None and str(v).strip():
                vals.append(f"[{r},{c}]={v}")
        if vals:
            print("  " + " | ".join(vals))
    wb.close()
