import sys
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

f = r"C:\Users\Hamed Ali Khan\Desktop\SQL tracking report.xlsx"
wb = load_workbook(f, data_only=True)
print("sheets:", wb.sheetnames)
for sn in wb.sheetnames:
    ws = wb[sn]
    print("=" * 110)
    print("SHEET:", sn, "| dims:", ws.max_row, "x", ws.max_column)
    for r in range(1, min(ws.max_row + 1, 60)):
        vals = []
        for c in range(1, min(ws.max_column + 1, 20)):
            v = ws.cell(r, c).value
            vals.append(str(v).replace("\xa0", " ").strip()[:28] if v is not None else "")
        if any(vals):
            print(" ", r, vals)
    if ws.max_row > 60:
        print("  ... (", ws.max_row - 60, "more rows )")
wb.close()
