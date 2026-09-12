import sys
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

f = r"C:\Users\Hamed Ali Khan\Desktop\SQL tracking report.xlsx"
wb = load_workbook(f, data_only=True)
for sn in wb.sheetnames:
    ws = wb[sn]
    print("=" * 130)
    print("SHEET:", sn, "| dims:", ws.max_row, "x", ws.max_column)
    # header row (row 6) full
    hdr = [str(ws.cell(6, c).value or "").replace("\xa0", " ").strip() for c in range(1, ws.max_column + 1)]
    print("HEADER:", hdr)
    # find any cell containing 'as per' anywhere
    hits = []
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if v is not None and "as per" in str(v).lower():
                hits.append((r, c, str(v)[:60]))
    print("'as per' cells:", hits[:20])
    # last 12 rows full
    print("--- last rows ---")
    for r in range(max(1, ws.max_row - 11), ws.max_row + 1):
        vals = []
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            vals.append(str(v).replace("\xa0", " ").strip()[:24] if v is not None else "")
        if any(vals):
            print(" ", r, vals)
wb.close()
