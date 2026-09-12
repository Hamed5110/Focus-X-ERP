import sys
from collections import defaultdict
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

path = r"C:\Users\Hamed Ali Khan\Downloads\Monthly_Commission_Atlas363679_ATLAS_ALUMINUM_W_L_L_.xlsx"
wb = load_workbook(path, data_only=True)
ws = wb[wb.sheetnames[0]]

# Reconstruct from shifted layout:
# Particulars = group label (dept/month/salesman/customer)
# Customer Name col = Payment Code
# Payment Code col = Collection
# Collection col = Contract
# Contract col = Eligible
# Eligible col = Team Eligible (repeated)
# last col = packed iDate
details = []
for i, r in enumerate(ws.iter_rows(values_only=True), 1):
    if i < 10:
        continue
    particulars = r[0]
    pay = r[1]
    coll = r[2]
    contract = r[3]
    elig = r[4]
    team_elig = r[5]
    if not isinstance(pay, str) or not pay.startswith("CI"):
        continue
    details.append({
        "row": i,
        "customer": str(particulars).strip() if particulars else "",
        "pay": pay,
        "coll": float(coll or 0),
        "contract": float(contract or 0),
        "elig": float(elig or 0),
        "team": float(team_elig or 0),
        "idate": r[10],
    })

print(f"detail rows (payment code in col B): {len(details)}")
print(f"sum coll={sum(d['coll'] for d in details):.2f}")
print(f"sum elig={sum(d['elig'] for d in details):.2f}")
print(f"first-pay elig={sum(d['elig'] for d in details if 'First Payment' in d['pay']):.2f}")
print(f"second-pay elig={sum(d['elig'] for d in details if 'Second Payment' in d['pay']):.2f}")
print(f"other-pay elig={sum(d['elig'] for d in details if 'First' not in d['pay'] and 'Second' not in d['pay']):.2f}")

by_pay = defaultdict(lambda: [0, 0, 0])
for d in details:
    by_pay[d["pay"]][0] += 1
    by_pay[d["pay"]][1] += d["coll"]
    by_pay[d["pay"]][2] += d["elig"]
print("\nby payment code (n / coll / elig)")
for k, v in sorted(by_pay.items()):
    print(f"  {k}: n={v[0]} coll={v[1]:.2f} elig={v[2]:.2f}")

print("\nheader totals row7:", [ws.cell(7, c).value for c in range(1, 12)])
print("row8 month:", [ws.cell(8, c).value for c in range(1, 12)])
wb.close()
