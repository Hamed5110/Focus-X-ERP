import sys
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

CUBE = r"C:\Users\Hamed Ali Khan\Downloads\Trade_Receivables_180_821618_Project_Tracking_III_Report_Atlas821618_ATLAS_ALUMINUM_W_L_L_.xlsx"

MIXED = [
    "Mr. Ahmed Abdulla Jaffar",
    "Mr. Ebrahim Ahmed",
    "Mr. Ebrahim Habib",
    "Mr. Hassan Abdul Amir",
    "Mr. Hassan Ali",
    "Mr. Hussain Ali",
    "Mr. Hussain Alsaleem",
    "Mr. Khalil Ebrahim",
]

wb = load_workbook(CUBE, data_only=True)
ws = wb.active
hr = None
for r in range(1, 12):
    vals = [str(ws.cell(r, c).value or "").replace("\xa0", " ").strip() for c in range(1, ws.max_column + 1)]
    if "Total Contract Amount" in vals:
        hr = r
        headers = vals
        break
idx = {h: i for i, h in enumerate(headers, 1)}
cube = {}
for r in range(hr + 1, ws.max_row + 1):
    name = str(ws.cell(r, idx["Particulars"]).value or "").strip()
    if not name or "grand" in name.lower() or name.lower() == "total":
        continue
    cube[name] = {"code": str(ws.cell(r, idx["Code"]).value or "").strip(),
                  "c": ws.cell(r, idx["Total Contract Amount"]).value,
                  "a": ws.cell(r, idx["Adv. Rct Amount"]).value}
wb.close()

print("Mixed-status name-groups: cube membership")
for n in MIXED:
    hit = cube.get(n)
    if hit:
        print(f"  IN CUBE : {n:35} code={hit['code']:>10} C={hit['c']} A={hit['a']}")
    else:
        print(f"  EXCLUDED: {n}")
