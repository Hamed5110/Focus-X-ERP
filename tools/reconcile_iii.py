# -*- coding: utf-8 -*-
"""Reconcile cube Excel (from Trade Receivables) vs Focus Query SQL Excel."""
import sys
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8", errors="replace")


def num(v):
    if v is None or v is False:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "")
    try:
        return float(s)
    except ValueError:
        return 0.0


def load(path):
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    headers = [
        str(ws.cell(6, c).value or "").replace("\xa0", " ").strip()
        for c in range(1, ws.max_column + 1)
    ]
    idx = {h: i for i, h in enumerate(headers, 1)}
    data = {}
    footer = None
    for r in range(7, ws.max_row + 1):
        name = str(ws.cell(r, 1).value or "").strip()
        code = str(ws.cell(r, idx["Code"]).value or "").strip() if "Code" in idx else ""
        rec = {
            "name": name,
            "c": num(ws.cell(r, idx["Total Contract Amount"]).value),
            "a": num(ws.cell(r, idx["Adv. Rct Amount"]).value),
        }
        if "grand" in name.lower():
            footer = rec
            continue
        if code:
            data[code] = rec
    wb.close()
    return data, footer


def main(cube_path, sql_path):
    cube, cf = load(cube_path)
    sql, sf = load(sql_path)
    common = set(cube) & set(sql)
    mc = sum(1 for k in common if abs(sql[k]["c"] - cube[k]["c"]) > 0.01)
    ma = sum(1 for k in common if abs(sql[k]["a"] - cube[k]["a"]) > 0.01)
    print("cube rows", len(cube), "sql rows", len(sql), "common", len(common))
    print("common Contract fails", mc, "Adv fails", ma)
    print("only sql", sorted(set(sql) - set(cube)))
    print("only cube", sorted(set(cube) - set(sql)))
    print("sum common cube", sum(cube[k]["c"] for k in common), sum(cube[k]["a"] for k in common))
    print("sum common sql ", sum(sql[k]["c"] for k in common), sum(sql[k]["a"] for k in common))
    print("sum all sql    ", sum(v["c"] for v in sql.values()), sum(v["a"] for v in sql.values()))
    print("cube footer    ", cf)
    print("sql footer     ", sf)
    if cf:
        print(
            "FOOTER TRAP Contract",
            cf["c"] - sum(v["c"] for v in cube.values()),
            "(cube footer minus cube row sum)",
        )


if __name__ == "__main__":
    cube = sys.argv[1] if len(sys.argv) > 1 else (
        r"C:\Users\Hamed Ali Khan\Downloads\Trade_Receivables_180_527369_"
        r"Project_Tracking_III_Report_Atlas527369_ATLAS_ALUMINUM_W_L_L_.xlsx"
    )
    sql = sys.argv[2] if len(sys.argv) > 2 else (
        r"C:\Users\Hamed Ali Khan\Downloads\III_prjoect_SQL352277_ATLAS_ALUMINUM_W_L_L_.xlsx"
    )
    main(cube, sql)
