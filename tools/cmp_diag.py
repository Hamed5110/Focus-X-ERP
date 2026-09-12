import sys
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

SQLX = r"C:\Users\Hamed Ali Khan\Downloads\iii_Project_Summ_SQL_report531004_ATLAS_ALUMINUM_W_L_L_.xlsx"
CUBE = r"C:\Users\Hamed Ali Khan\Downloads\Trade_Receivables_180_821618_Project_Tracking_III_Report_Atlas821618_ATLAS_ALUMINUM_W_L_L_.xlsx"


def num(v):
    if v is None or v is False:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "")
    if s in ("", "-", "None", "False"):
        return 0.0
    try:
        return float(s)
    except ValueError:
        return 0.0


def find_header_row(ws, must_have):
    for r in range(1, 12):
        vals = [str(ws.cell(r, c).value or "").replace("\xa0", " ").strip() for c in range(1, ws.max_column + 1)]
        if all(any(mh == v for v in vals) for mh in must_have):
            return r, vals
    return None, None


# ---- SQL diagnostic export ----
wb = load_workbook(SQLX, data_only=True)
ws = wb.active
hr, headers = find_header_row(ws, ["Report Status", "No. of Accounts"])
print("SQL header row:", hr, headers)
idx = {h: i for i, h in enumerate(headers, 1)}
sql_rows = {}
for r in range(hr + 1, ws.max_row + 1):
    name = str(ws.cell(r, idx["Report Status"]).value or "").strip()
    if not name or "grand" in name.lower() or name.lower() == "total":
        continue
    mid = num(ws.cell(r, idx["No. of Accounts"]).value)
    sql_rows[name] = {
        "mid": int(mid),
        "c": num(ws.cell(r, idx["Total Contract Amount"]).value),
        "a": num(ws.cell(r, idx["Adv. Rct Amount"]).value),
        "b": num(ws.cell(r, idx["Balance Amount"]).value),
        "p": num(ws.cell(r, idx["Plan Value"]).value),
    }
wb.close()
print("SQL rows:", len(sql_rows))

# ---- Cube export ----
wb = load_workbook(CUBE, data_only=True)
ws = wb.active
hr, headers = find_header_row(ws, ["Total Contract Amount"])
print("\nCUBE header row:", hr, headers)
idx = {h: i for i, h in enumerate(headers, 1)}
name_col = 1
for cand in ("Particulars", "Name", "Account Name"):
    if cand in idx:
        name_col = idx[cand]
        break
cube_rows = {}
for r in range(hr + 1, ws.max_row + 1):
    name = str(ws.cell(r, name_col).value or "").strip()
    if not name or "grand" in name.lower() or name.lower() == "total":
        continue
    cube_rows[name] = {
        "c": num(ws.cell(r, idx["Total Contract Amount"]).value),
        "a": num(ws.cell(r, idx["Adv. Rct Amount"]).value),
    }
wb.close()
print("CUBE rows:", len(cube_rows))

cs, ss = set(cube_rows), set(sql_rows)
print("\n==== SQL-only accounts (in SQL, not in cube) ====")
for n in sorted(ss - cs):
    v = sql_rows[n]
    print(f"  id={v['mid']:<7} {n[:55]:57} C={v['c']:>10.2f} A={v['a']:>10.2f} B={v['b']:>10.2f} P={v['p']:>10.2f}")
print("\n==== Cube-only accounts (in cube, not in SQL) ====")
for n in sorted(cs - ss):
    v = cube_rows[n]
    print(f"  {n[:60]:62} C={v['c']:>10.2f} A={v['a']:>10.2f}")

print("\nSQL totals: C={:.2f} A={:.2f}".format(sum(v["c"] for v in sql_rows.values()), sum(v["a"] for v in sql_rows.values())))
print("CUBE totals: C={:.2f} A={:.2f}".format(sum(v["c"] for v in cube_rows.values()), sum(v["a"] for v in cube_rows.values())))
