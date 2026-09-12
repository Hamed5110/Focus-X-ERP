import sys
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

DL = r"C:\Users\Hamed Ali Khan\Downloads"

def read_rows(path, sheet=None):
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
    rows = []
    for r in ws.iter_rows(values_only=True):
        rows.append([("" if v is None else v) for v in r])
    wb.close()
    return rows

def num(v):
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace(",", "").replace("\xa0", "").strip()
    if s in ("", "-"):
        return 0.0
    try:
        return float(s)
    except ValueError:
        return 0.0

# ---------- SQL summary export ----------
sqlf = DL + r"\iii_Project_Summ_SQL_report324657_ATLAS_ALUMINUM_W_L_L_.xlsx"
rows = read_rows(sqlf)
print("=== SQL summary export ===")
sql = {}
for r in rows:
    line = [str(x).strip() for x in r]
    joined = " ".join(line)
    for key, sid in (("1 Pending - I", 1), ("2 In Progress - II", 2), ("3 Partial Consumed - III", 3)):
        if key in joined:
            vals = [num(x) for x in r]
            vals = [v for v in vals if v != 0 or True]
            # find numeric cells: contract, adv, balance, plan, count
            nums = [num(x) for x in r if str(x).strip() != "" and isinstance(x, (int, float))]
            sql[sid] = nums
            print(key, "->", nums)

# ---------- cube exports ----------
cubes = {
    1: DL + r"\Trade_Receivables_180_643162_Project_Tracking_I_Report_Atlas643162_ATLAS_ALUMINUM_W_L_L_.xlsx",
    2: DL + r"\Trade_Receivables_180_264376_Project_Tracking_II_Report_Atlas264376_ATLAS_ALUMINUM_W_L_L_.xlsx",
    3: DL + r"\Trade_Receivables_180_145256_Project_Tracking_III_Report_Atlas145256_ATLAS_ALUMINUM_W_L_L_.xlsx",
}
print("\n=== Cube exports (row sums) ===")
cube = {}
for sid, path in cubes.items():
    rows = read_rows(path)
    # find header row containing 'Particulars'
    hi = None
    for i, r in enumerate(rows):
        if str(r[0]).strip() == "Particulars":
            hi = i
            break
    hdr = [str(x).strip() for x in rows[hi]]
    ci_c = hdr.index("Total Contract Amount")
    ci_a = hdr.index("Adv. Rct Amount")
    tc = ta = 0.0
    n = 0
    for r in rows[hi + 1:]:
        name = str(r[0]).strip()
        if name in ("", "Grand Total"):
            continue
        tc += num(r[ci_c])
        ta += num(r[ci_a])
        n += 1
    cube[sid] = (tc, ta, n)
    print(f"status {sid}: rows={n} contract={tc:,.2f} adv={ta:,.2f}")

print("\n=== DIFF (SQL - cube) ===")
for sid in (1, 2, 3):
    s = sql.get(sid, [])
    c = cube.get(sid)
    if not s or not c:
        continue
    # sql nums order: contract, adv, balance, plan, count (may include row index)
    print(f"status {sid}: sql nums={s}")
    print(f"          cube: contract={c[0]:,.2f} adv={c[1]:,.2f} rows={c[2]}")
