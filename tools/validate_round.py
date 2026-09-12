# -*- coding: utf-8 -*-
import csv
import sys
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8", errors="replace")


def num(v):
    if v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", ""))
    except Exception:
        return 0.0


wb = load_workbook(
    r"C:\Users\Hamed Ali Khan\Downloads\Trade_Receivables_180_931647_Project_Tracking_III_Report_Atlas931647_ATLAS_ALUMINUM_W_L_L_.xlsx",
    data_only=True,
)
ws = wb.active
headers = [
    str(ws.cell(6, c).value or "").replace("\xa0", " ").strip()
    for c in range(1, ws.max_column + 1)
]
idx = {h: i for i, h in enumerate(headers, 1)}
cube = {}
for r in range(7, ws.max_row + 1):
    name = str(ws.cell(r, 1).value or "").strip()
    if "grand" in name.lower():
        continue
    code = str(ws.cell(r, idx["Code"]).value or "").strip()
    if not code:
        continue
    cube[code] = {
        "c": num(ws.cell(r, idx["Total Contract Amount"]).value),
        "a": num(ws.cell(r, idx["Adv. Rct Amount"]).value),
    }
wb.close()

sql = {}
with open(
    r"C:\Users\Hamed Ali Khan\Focus-X-ERP\tools\sql_round_amounts.csv",
    encoding="utf-8-sig",
) as f:
    for row in csv.DictReader(f):
        sql[row["sCode"]] = {
            "c": num(row["ContractRaw"]),
            "cr": num(row["ContractRoundV"]),
            "a": num(row["AdvRaw"]),
            "ar": num(row["AdvRoundVoucher"]),
        }

print("AdvRoundVoucher vs cube Adv")
fail = []
for k in sorted(set(cube) & set(sql)):
    if abs(sql[k]["ar"] - cube[k]["a"]) > 0.01:
        fail.append((k, cube[k]["a"], sql[k]["a"], sql[k]["ar"]))
print("fails", len(fail))
for x in fail[:40]:
    print(x)

print("\nROUND(ContractRaw,0) vs cube")
failc = []
for k in sorted(set(cube) & set(sql)):
    if abs(round(sql[k]["c"]) - cube[k]["c"]) > 0.01:
        failc.append((k, cube[k]["c"], sql[k]["c"], round(sql[k]["c"])))
print("fails", len(failc))
for x in failc[:20]:
    print(x)

print("\nContractRoundV vs cube")
failcv = []
for k in sorted(set(cube) & set(sql)):
    if abs(sql[k]["cr"] - cube[k]["c"]) > 0.01:
        failcv.append((k, cube[k]["c"], sql[k]["c"], sql[k]["cr"]))
print("fails", len(failcv))
for x in failcv[:20]:
    print(x)

common = set(cube) & set(sql)
print("\nsums cube", sum(v["c"] for v in cube.values()), sum(v["a"] for v in cube.values()))
print("sums AdvRoundVoucher", sum(sql[k]["ar"] for k in common))
print("sums ROUND contract", sum(round(sql[k]["c"]) for k in common))
