import sys
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

CUBE = r"C:\Users\Hamed Ali Khan\Downloads\Trade_Receivables_180_821618_Project_Tracking_III_Report_Atlas821618_ATLAS_ALUMINUM_W_L_L_.xlsx"

wb = load_workbook(CUBE, data_only=True)
ws = wb.active
hr = None
for r in range(1, 12):
    vals = [str(ws.cell(r, c).value or "").replace("\xa0", " ").strip() for c in range(1, ws.max_column + 1)]
    if "Total Contract Amount" in vals:
        hr = r
        headers = vals
        break
idx = {h: i for i, h in enumerate(headers, 1)}

print("Searching cube export for codes AC-3758 / AC-5850 and the 2 names...")
found = False
for r in range(hr + 1, ws.max_row + 1):
    code = str(ws.cell(r, idx["Code"]).value or "").strip()
    name = str(ws.cell(r, idx["Particulars"]).value or "").strip()
    if code in ("AC-3758", "AC-5850") or name in ("Mr. Ahmed Abdulla Jaffar", "Mr. Hussain Ali"):
        row = {h: ws.cell(r, i).value for h, i in idx.items()}
        print("FOUND:", row)
        found = True
if not found:
    print("NOT FOUND in cube export - the cube truly excludes them.")

# Also dump distinct Site Status / Pipeline / Advance Payment Status values across cube rows
for col in ("Site Status", "Pipeline", "Advance Payment Status", "Final Measurement Status", "SJO processing"):
    if col in idx:
        vals = {}
        for r in range(hr + 1, ws.max_row + 1):
            v = str(ws.cell(r, idx[col]).value or "").strip()
            vals[v] = vals.get(v, 0) + 1
        print(f"\nDistinct '{col}' values:", vals)
wb.close()
