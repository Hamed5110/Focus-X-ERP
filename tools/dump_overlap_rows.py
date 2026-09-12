import sys, glob
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

base = r"C:\Users\Hamed Ali Khan\Downloads"
files = {
    "I  ": glob.glob(base + r"\*Project_Tracking_I_Report_Atlas362696*.xlsx")[0],
    "II ": glob.glob(base + r"\*Project_Tracking_II_Report_Atlas314796*.xlsx")[0],
    "III": glob.glob(base + r"\*Project_Tracking_III_Report_Atlas821618*.xlsx")[0],
}

targets = ["Ahmed Isa Marhoon", "Alshakhs", "Alhujairi", "Mahdi Jaafar"]


def load(path):
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    hr = None
    headers = None
    for r in range(1, 12):
        vals = [str(ws.cell(r, c).value or "").replace("\xa0", " ").strip() for c in range(1, ws.max_column + 1)]
        if "Total Contract Amount" in vals:
            hr = r
            headers = vals
            break
    idx = {h: i for i, h in enumerate(headers, 1) if h}
    rows = {}
    for r in range(hr + 1, ws.max_row + 1):
        n = str(ws.cell(r, idx["Particulars"]).value or "").strip()
        if n and "grand" not in n.lower() and n.lower() != "total":
            rows[n] = {h: ws.cell(r, i).value for h, i in idx.items()}
    wb.close()
    return headers, rows


for tag, f in files.items():
    headers, rows = load(f)
    print("=" * 100)
    print(tag, "report columns:", [h for h in headers if h])
    for name, row in rows.items():
        if any(t.lower() in name.lower() for t in targets):
            print()
            print(" ", name)
            for h, v in row.items():
                if v not in (None, "", 0, 0.0, False):
                    print("    ", h, "=", v)
