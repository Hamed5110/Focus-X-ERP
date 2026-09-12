import sys
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

files = [
    r"C:\Commission reports\Aknan Commission - from April to July.xlsx",
    r"C:\Commission reports\Atlas Aug 2026.xlsx",
    r"C:\Commission reports\Rabab - Payment.xlsx",
]

for path in files:
    print("=" * 80)
    print(path)
    wb = load_workbook(path, data_only=True)
    print("sheets:", wb.sheetnames)
    for sn in wb.sheetnames:
        ws = wb[sn]
        print(f"\n--- sheet: {sn}  dims={ws.dimensions} max_row={ws.max_row} max_col={ws.max_column} ---")
        for i, r in enumerate(ws.iter_rows(values_only=True), 1):
            if any(v is not None and str(v).strip() != "" for v in r):
                cells = [("" if v is None else str(v).replace("\n", " | "))[:50] for v in r]
                # trim trailing empties
                while cells and cells[-1] == "":
                    cells.pop()
                print(f"{i:03d}| " + " || ".join(cells))
    wb.close()
    print()
