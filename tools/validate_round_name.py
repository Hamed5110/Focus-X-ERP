# -*- coding: utf-8 -*-
import csv
import sys
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8", errors="replace")


def num(v):
    if v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", ""))
    except Exception:
        return 0.0


wb = load_workbook(
    r"C:\Users\Hamed Ali Khan\Downloads\Trade_Receivables_180_931647_Project_Tracking_III_Report_Atlas931647_ATLAS_ALUMINUM_W_L_L_.xlsx",
    data_only=True,
)
ws = wb.active
headers = [
    str(ws.cell(6, c).value or "").replace("\xa0", " ").strip()
    for c in range(1, ws.max_column + 1)
]
idx = {h: i for i, h in enumerate(headers, 1)}
cube = {}
for r in range(7, ws.max_row + 1):
    name = str(ws.cell(r, 1).value or "").strip()
    if "grand" in name.lower():
        continue
    code = str(ws.cell(r, idx["Code"]).value or "").strip()
    if not code:
        continue
    cube[code] = {
        "c": num(ws.cell(r, idx["Total Contract Amount"]).value),
        "a": num(ws.cell(r, idx["Adv. Rct Amount"]).value),
    }
wb.close()

sql = {}
with open(
    r"C:\Users\Hamed Ali Khan\Focus-X-ERP\tools\sql_round_name.csv",
    encoding="utf-8-sig",
) as f:
    for row in csv.DictReader(f):
        sql[row["sCode"]] = {k: num(row[k]) for k in row if k != "sCode"}

common = set(cube) & set(sql)
print("codes cube", len(cube), "sql", len(sql), "common", len(common))


def show(label, key):
    fails = []
    for c in sorted(common):
        if abs(sql[c][key] - cube[c]["a" if "Adv" in key or key.startswith("Adv") else "c"]) > 0.01:
            # pick cube field
            pass
    return fails


tests = [
    ("ContractRoundT", "c", "ContractRoundT"),
    ("ContractRoundV", "c", "ContractRoundV"),
    ("AdvRoundT", "a", "AdvRoundT"),
    ("AdvRoundV", "a", "AdvRoundV"),
]
for label, cube_field, sql_field in tests:
    fails = []
    for c in sorted(common):
        if abs(sql[c][sql_field] - cube[c][cube_field]) > 0.01:
            fails.append(
                (
                    c,
                    cube[c][cube_field],
                    sql[c]["ContractRaw" if cube_field == "c" else "AdvRaw"],
                    sql[c][sql_field],
                )
            )
    print(f"\n{label} fails {len(fails)}")
    for x in fails[:15]:
        print(x)

print("\nsums cube", sum(v["c"] for v in cube.values()), sum(v["a"] for v in cube.values()))
print("ContractRoundT", sum(sql[c]["ContractRoundT"] for c in common))
print("AdvRoundT", sum(sql[c]["AdvRoundT"] for c in common))
print("AdvRoundV", sum(sql[c]["AdvRoundV"] for c in common))
