import sys
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

FILES = {
    "III fresh 821618": r"C:\Users\Hamed Ali Khan\Downloads\Trade_Receivables_180_821618_Project_Tracking_III_Report_Atlas821618_ATLAS_ALUMINUM_W_L_L_.xlsx",
    "III older 60810": r"C:\Users\Hamed Ali Khan\Downloads\Trade_Receivables_180_60810_Project_Tracking_III_Report_Atlas60810_ATLAS_ALUMINUM_W_L_L_.xlsx",
    "III oldest 931647": r"C:\Users\Hamed Ali Khan\Downloads\Trade_Receivables_180_931647_Project_Tracking_III_Report_Atlas931647_ATLAS_ALUMINUM_W_L_L_.xlsx",
    "III mid 487465": r"C:\Users\Hamed Ali Khan\Downloads\Trade_Receivables_180_487465_Project_Tracking_III_Report_Atlas487465_ATLAS_ALUMINUM_W_L_L_.xlsx",
}

PROBES = ["Abdulrahman Abdullah", "Mr. Ahmed Darraj", "Sq-atl-924 Sayed Mustafa Hasan Mohamed Ali",
          "Sq-atl-939 Khaled Mohamed Rashed Aldoseri", "Sayed Noaman Ismaeel Ebrahim Mahfood",
          "Mr. Ahmed Abdulla Jaffar", "Mr. Hussain Ali"]


def num(v):
    if v is None or v is False:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "")
    try:
        return float(s)
    except ValueError:
        return 0.0


for label, path in FILES.items():
    try:
        wb = load_workbook(path, data_only=True)
    except Exception as e:
        print(f"{label}: cannot open ({e})")
        continue
    ws = wb.active
    hr = None
    for r in range(1, 12):
        vals = [str(ws.cell(r, c).value or "").replace("\xa0", " ").strip() for c in range(1, ws.max_column + 1)]
        if "Total Contract Amount" in vals:
            hr = r
            headers = vals
            break
    idx = {h: i for i, h in enumerate(headers, 1)}
    ncol = idx.get("Particulars", 1)
    rows = {}
    for r in range(hr + 1, ws.max_row + 1):
        n = str(ws.cell(r, ncol).value or "").strip()
        if n and "grand" not in n.lower() and n.lower() != "total":
            rows[n] = (num(ws.cell(r, idx["Total Contract Amount"]).value),
                       num(ws.cell(r, idx["Adv. Rct Amount"]).value))
    wb.close()
    print(f"\n== {label} ({len(rows)} rows) ==")
    for p in PROBES:
        if p in rows:
            print(f"  {p[:50]:52} C={rows[p][0]:>10.2f} A={rows[p][1]:>10.2f}")
        else:
            print(f"  {p[:50]:52}  -- NOT IN REPORT --")
