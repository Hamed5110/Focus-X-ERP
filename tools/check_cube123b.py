import sys
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

base = r"C:\Users\Hamed Ali Khan\Downloads"
files = {
    "I": base + r"\Trade_Receivables_180_362696_Project_Tracking_I_Report_Atlas362696_ATLAS_ALUMINUM_W_L_L_.xlsx",
    "II": base + r"\Trade_Receivables_180_314796_Project_Tracking_II_Report_Atlas314796_ATLAS_ALUMINUM_W_L_L_.xlsx",
    "III": base + r"\Trade_Receivables_180_821618_Project_Tracking_III_Report_Atlas821618_ATLAS_ALUMINUM_W_L_L_.xlsx",
}


def load(path):
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    hr = None
    headers = None
    for r in range(1, 12):
        vals = [str(ws.cell(r, c).value or "").replace("\xa0", " ").strip() for c in range(1, ws.max_column + 1)]
        if "Total Contract Amount" in vals:
            hr = r
            headers = vals
            break
    idx = {h: i for i, h in enumerate(headers, 1) if h}
    rows = {}
    for r in range(hr + 1, ws.max_row + 1):
        n = str(ws.cell(r, idx["Particulars"]).value or "").replace("\xa0", " ").strip()
        if n and "grand" not in n.lower() and n.lower() != "total":
            rows[n] = {h: ws.cell(r, i).value for h, i in idx.items()}
    wb.close()
    return rows


data = {k: load(v) for k, v in files.items()}
for k, rows in data.items():
    tc = sum(float(r.get("Total Contract Amount") or 0) for r in rows.values())
    ta = sum(float(r.get("Adv. Rct Amount") or 0) for r in rows.values())
    print(f"Cube {k}: {len(rows)} rows | Contract={tc:,.0f} | Adv={ta:,.0f}")

s1, s2, s3 = set(data["I"]), set(data["II"]), set(data["III"])
print()
print("I and III both:", sorted(s1 & s3))
print("II and III both:", sorted(s2 & s3))
print("I and II both:", sorted(s1 & s2))
print("I and II and III:", sorted(s1 & s2 & s3))
