import sys, csv
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

base = r"C:\Users\Hamed Ali Khan\Downloads"
cube_files = {
    1: base + r"\Trade_Receivables_180_362696_Project_Tracking_I_Report_Atlas362696_ATLAS_ALUMINUM_W_L_L_.xlsx",
    2: base + r"\Trade_Receivables_180_314796_Project_Tracking_II_Report_Atlas314796_ATLAS_ALUMINUM_W_L_L_.xlsx",
    3: base + r"\Trade_Receivables_180_821618_Project_Tracking_III_Report_Atlas821618_ATLAS_ALUMINUM_W_L_L_.xlsx",
}


def load_cube(path):
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    hr = None
    headers = None
    for r in range(1, 12):
        vals = [str(ws.cell(r, c).value or "").replace("\xa0", " ").strip() for c in range(1, ws.max_column + 1)]
        if "Total Contract Amount" in vals:
            hr = r
            headers = vals
            break
    idx = {h: i for i, h in enumerate(headers, 1) if h}
    rows = {}
    for r in range(hr + 1, ws.max_row + 1):
        n = str(ws.cell(r, idx["Particulars"]).value or "").replace("\xa0", " ").strip()
        if n and "grand" not in n.lower() and n.lower() != "total":
            rows[n] = {
                "C": float(ws.cell(r, idx["Total Contract Amount"]).value or 0),
                "A": float(ws.cell(r, idx["Adv. Rct Amount"]).value or 0),
                "P": float(ws.cell(r, idx["Plan Value"]).value or 0) if "Plan Value" in idx else 0.0,
            }
    wb.close()
    return rows


# local account-level
local = {1: {}, 2: {}, 3: {}}
with open(r"C:\Users\Hamed Ali Khan\Focus-X-ERP\tools\acct_level_all.csv", encoding="utf-8-sig") as f:
    for row in csv.DictReader(f):
        st = int(row["ReportStatus"])
        name = row["Name"].strip()
        # name-level rollup (query is account-level; cube is name-level)
        cur = local[st].setdefault(name, {"C": 0.0, "A": 0.0, "P": 0.0})
        cur["C"] += float(row["Total Contract Amount"] or 0)
        cur["A"] += float(row["Adv. Rct Amount"] or 0)
        cur["P"] += float(row["Plan Value"] or 0)

for st in (1, 2, 3):
    cube = load_cube(cube_files[st])
    loc = local[st]
    print("=" * 90)
    print(f"STATUS {st}: local names={len(loc)} | cube rows={len(cube)}")
    only_loc = sorted(set(loc) - set(cube))
    only_cube = sorted(set(cube) - set(loc))
    print(f"  only in LOCAL ({len(only_loc)}):", only_loc)
    print(f"  only in CUBE  ({len(only_cube)}):", only_cube)
    amt_diffs = []
    for n in sorted(set(loc) & set(cube)):
        dc = loc[n]["C"] - cube[n]["C"]
        da = loc[n]["A"] - cube[n]["A"]
        dp = loc[n]["P"] - cube[n]["P"]
        if abs(dc) > 0.5 or abs(da) > 0.5 or abs(dp) > 0.5:
            amt_diffs.append((n, dc, da, dp))
    print(f"  amount diffs on common names ({len(amt_diffs)}):")
    for n, dc, da, dp in amt_diffs:
        print(f"    {n}: dC={dc:,.0f} dA={da:,.0f} dP={dp:,.0f}")
